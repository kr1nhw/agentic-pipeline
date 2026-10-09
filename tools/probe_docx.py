#!/usr/bin/env python3
"""What does an Office fixture actually contain, and what does the session see?

The pipeline reported "table candidate(s) detected but none extracted" with a
15-character text layer and zero blocks. Either the file is not a document (in
which case flagging it is right) or the reader is not seeing what python-docx
sees (in which case content is being dropped). This prints both views.

    tools/rpy.sh --docker tools/probe_docx.py /data/mm-lbfs/corpus/docx/<file>.docx
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path


def raw_xml_view(path: Path) -> dict:
    """What the OOXML package literally stores."""
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        body = ""
        for candidate in ("word/document.xml", "ppt/presentation.xml"):
            if candidate in names:
                body = archive.read(candidate).decode("utf-8", "replace")
                break
        return {
            "parts": len(names),
            "body_bytes": len(body),
            "w:tbl": body.count("<w:tbl>"),
            "w:tr": body.count("<w:tr>"),
            "w:tc": body.count("<w:tc>"),
            "w:p": body.count("<w:p "),
            "a:tbl": body.count("<a:tbl>"),
            "text": "".join(__import__("re").findall(r"<w:t[^>]*>([^<]*)</w:t>", body))[:200],
            "a_text": "".join(__import__("re").findall(r"<a:t>([^<]*)</a:t>", body))[:200],
        }


def python_docx_view(path: Path) -> dict:
    try:
        import docx
    except ImportError:
        return {"error": "python-docx not installed"}
    document = docx.Document(str(path))
    tables = []
    for table in document.tables:
        tables.append([[cell.text for cell in row.cells] for row in table.rows])
    nested = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                nested += len(cell.tables)
    return {
        "paragraphs": len(document.paragraphs),
        "paragraph_text": [p.text for p in document.paragraphs][:6],
        "tables": len(document.tables),
        "nested_tables": nested,
        "table_shapes": [[len(r), len(r[0]) if r else 0] for r in tables],
        "table_cells": [[[c[:24] for c in row] for row in rows][:4] for rows in tables][:2],
    }


def session_view(path: Path) -> dict:
    """What docpipe's own office session reports."""
    try:
        from docpipe.ingest import office
        from docpipe.config import Config
    except ImportError as exc:
        return {"error": f"docpipe unavailable: {exc}"}
    session = office.open_docx(path, "probe", Config())
    out: dict = {"type": type(session).__name__}
    for attr in ("sections", "section_text", "flow_text", "n_sections"):
        if hasattr(session, attr):
            value = getattr(session, attr)
            value = value() if callable(value) else value
            out[attr] = json.dumps(value, default=str)[:600]
    return out


def main(path_str: str) -> int:
    path = Path(path_str)
    print(f"file      {path}  ({path.stat().st_size} bytes)")
    print("raw       " + json.dumps(raw_xml_view(path), ensure_ascii=False))
    print("docx      " + json.dumps(python_docx_view(path), ensure_ascii=False)[:900])
    print("session   " + json.dumps(session_view(path), ensure_ascii=False)[:900])
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1]))
